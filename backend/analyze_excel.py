import openpyxl
import sys

# Načti Excel soubor
excel_path = r"C:\Users\martin.korcak\Downloads\Tabulka tržeb.xlsx"

try:
    wb = openpyxl.load_workbook(excel_path)

    print("=" * 60)
    print("ANALÝZA EXCEL SOUBORU")
    print("=" * 60)
    print(f"\nSoubor: {excel_path}")
    print(f"\nPočet listů: {len(wb.sheetnames)}")
    print(f"Názvy listů: {wb.sheetnames}")

    # Projdi každý list
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n{'=' * 60}")
        print(f"LIST: {sheet_name}")
        print(f"{'=' * 60}")
        print(f"Rozměry: {ws.max_row} řádků × {ws.max_column} sloupců")

        # Zobraz VŠECHNY sloupce prvního řádku (hlavičky)
        print(f"\nHLAVIČKY (řádek 1) - všech {ws.max_column} sloupců:")
        print("-" * 100)
        header_row = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            value = cell.value if cell.value is not None else ""
            header_row.append(str(value)[:15])

        # Vypiš po 10
        for i in range(0, len(header_row), 10):
            chunk = header_row[i:i+10]
            print(f"Sloupce {i+1}-{min(i+10, len(header_row))}: {' | '.join(chunk)}")

        # Zobraz prvních 5 řádků dat
        print(f"\nPrvních 5 řádků dat (všechny sloupce):")
        print("-" * 100)

        for row_idx in range(2, min(7, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                value = cell.value
                if value is None:
                    value = ""
                # Ořízni dlouhé hodnoty
                value_str = str(value)[:15]
                row_data.append(value_str)

            # Vypiš po 10
            print(f"\nŘádek {row_idx}:")
            for i in range(0, len(row_data), 10):
                chunk = row_data[i:i+10]
                print(f"  Sl. {i+1}-{min(i+10, len(row_data))}: {' | '.join(chunk)}")

        if ws.max_row > 10:
            print(f"\n... (celkem {ws.max_row} řádků)")

    wb.close()
    print("\n" + "=" * 60)
    print("ANALÝZA DOKONČENA")
    print("=" * 60)

except FileNotFoundError:
    print(f"CHYBA: Soubor nenalezen: {excel_path}")
    print("Zkontrolujte prosím cestu k souboru.")
except Exception as e:
    print(f"CHYBA při čtení souboru: {str(e)}")
    import traceback
    traceback.print_exc()
