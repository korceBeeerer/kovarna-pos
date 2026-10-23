import openpyxl
from datetime import datetime
import json
import os
import shutil
from werkzeug.utils import secure_filename

# KONFIGURACE
# Excel soubor se nahrává z webového rozhraní (⚙️ Nastavení) do složky data/
# a cesta k němu se ukládá do config.json vedle tohoto souboru.
BACKEND_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(BACKEND_DIR, 'config.json')
DATA_DIR = os.path.join(os.path.dirname(BACKEND_DIR), 'data')
DEFAULT_EXCEL_FILE = r"C:\Users\martin.korcak\Downloads\Tabulka tržeb.xlsx"
WORKSHEET_NAME = "Tržby"

# Konstanty pro strukturu Excelu
ROWS_PER_DAY = 50  # Každý den zabírá 50 řádků
HEADER_ROW_OFFSET = 0  # Hlavička je na prvním řádku bloku
PRODUCTS_START_OFFSET = 1  # Produkty začínají na druhém řádku bloku

# Sloupce
COL_NUMBER = 1  # A
COL_PRODUCT = 2  # B
COL_PRICE = 3  # C
COL_QUANTITY = 4  # D
COL_TOTAL = 5  # E


def load_config():
    """Načte konfiguraci z config.json (pokud existuje)."""
    try:
        if os.path.exists(CONFIG_FILE):
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
    except Exception as e:
        print(f"Chyba při načítání konfigurace: {str(e)}")
    return {}


EXCEL_FILE = load_config().get('excel_file') or DEFAULT_EXCEL_FILE


def get_config():
    """Vrátí aktuální konfiguraci pro webové rozhraní."""
    return {
        "excel_file": EXCEL_FILE,
        "excel_filename": os.path.basename(EXCEL_FILE),
        "exists": os.path.exists(EXCEL_FILE)
    }


def validate_excel_file(path):
    """
    Ověří, že soubor je Excel s listem WORKSHEET_NAME.
    Vrací text chyby, nebo None, pokud je soubor v pořádku.
    """
    if not os.path.isfile(path):
        return f"Soubor nenalezen: {path}"

    if not path.lower().endswith(('.xlsx', '.xlsm')):
        return "Soubor musí být ve formátu .xlsx nebo .xlsm"

    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        sheet_names = wb.sheetnames
        wb.close()
    except Exception as e:
        return f"Soubor nelze otevřít jako Excel: {str(e)}"

    if WORKSHEET_NAME not in sheet_names:
        return f"Soubor neobsahuje list '{WORKSHEET_NAME}' (nalezené listy: {', '.join(sheet_names)})"

    return None


def set_excel_file(path):
    """
    Ověří a nastaví novou cestu k Excel souboru a uloží ji do config.json.

    Returns:
        dict: {"success": bool, "message": str, "error": str (volitelné)}
    """
    global EXCEL_FILE

    path = (path or '').strip().strip('"').strip()
    if not path:
        return {"success": False, "error": "Cesta k souboru nesmí být prázdná"}

    error = validate_excel_file(path)
    if error:
        return {"success": False, "error": error}

    config = load_config()
    config['excel_file'] = path
    try:
        with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump(config, f, ensure_ascii=False, indent=2)
    except Exception as e:
        return {"success": False, "error": f"Nepodařilo se uložit konfiguraci: {str(e)}"}

    EXCEL_FILE = path
    print(f"Nastavena nová cesta k Excelu: {EXCEL_FILE}")
    return {"success": True, "message": "Cesta k Excel souboru byla uložena"}


def save_uploaded_excel(uploaded_file):
    """
    Uloží Excel nahraný z webového rozhraní do složky data/ a nastaví ho
    jako aktuální soubor. Pokud tam soubor se stejným názvem už je,
    nejdřív ho zazálohuje do data/backup/.

    Args:
        uploaded_file: werkzeug FileStorage z request.files

    Returns:
        dict: {"success": bool, "message": str, "error": str (volitelné)}
    """
    filename = secure_filename(uploaded_file.filename or '')
    if not filename.lower().endswith(('.xlsx', '.xlsm')):
        return {"success": False, "error": "Soubor musí být ve formátu .xlsx nebo .xlsm"}

    os.makedirs(DATA_DIR, exist_ok=True)
    target = os.path.join(DATA_DIR, filename)
    tmp_path = os.path.join(DATA_DIR, '_nahravani' + os.path.splitext(filename)[1])

    try:
        uploaded_file.save(tmp_path)

        error = validate_excel_file(tmp_path)
        if error:
            return {"success": False, "error": error}

        # Nepřepisuj existující soubor bez zálohy
        if os.path.exists(target):
            backup_dir = os.path.join(DATA_DIR, 'backup')
            os.makedirs(backup_dir, exist_ok=True)
            name, ext = os.path.splitext(filename)
            timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            backup_path = os.path.join(backup_dir, f"{name}_PREPSANO_{timestamp}{ext}")
            shutil.copy2(target, backup_path)
            print(f"Původní soubor zazálohován: {backup_path}")

        os.replace(tmp_path, target)

    except PermissionError:
        return {"success": False, "error": f"Soubor {filename} nelze přepsat - není otevřený v Excelu?"}
    except Exception as e:
        return {"success": False, "error": f"Chyba při ukládání souboru: {str(e)}"}
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

    result = set_excel_file(target)
    if result['success']:
        result['message'] = f"Soubor {filename} byl nahrán a nastaven"
    return result


def get_day_start_row(day):
    """
    Vypočítá řádek, kde začíná blok pro daný den.
    Den 1 -> řádek 1
    Den 2 -> řádek 51
    Den 3 -> řádek 101
    atd.
    """
    return 1 + (day - 1) * ROWS_PER_DAY


def load_products_from_excel():
    """
    Načte seznam produktů z prvního dne v Excelu.
    Vrací seznam produktů s cenami.
    """
    try:
        if not os.path.exists(EXCEL_FILE):
            print(f"VAROVÁNÍ: Excel soubor nenalezen: {EXCEL_FILE}")
            return []

        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb[WORKSHEET_NAME]

        products = []
        day_start = get_day_start_row(1)

        # Projdi všechny řádky prvního dne
        for offset in range(PRODUCTS_START_OFFSET, ROWS_PER_DAY):
            row_num = day_start + offset
            product_name = ws.cell(row=row_num, column=COL_PRODUCT).value
            price = ws.cell(row=row_num, column=COL_PRICE).value

            # Přidej pouze produkty s názvem a cenou
            if product_name and price:
                # Normalize název (odstranit bílé znaky)
                product_name = str(product_name).strip()
                if product_name:
                    products.append({
                        'id': len(products) + 1,
                        'name': product_name,
                        'price': float(price) if price else 0,
                        'row_offset': offset  # Relativní pozice v bloku (pro rychlé hledání)
                    })

        wb.close()
        print(f"Načteno {len(products)} produktů z Excelu")
        return products

    except Exception as e:
        print(f"Chyba při načítání produktů: {str(e)}")
        import traceback
        traceback.print_exc()
        return []


def find_product_row_in_day(ws, day_start_row, product_name):
    """
    Najde řádek pro daný produkt v konkrétním dni.
    Vrací číslo řádku nebo None, pokud produkt nebyl nalezen.
    """
    product_name_normalized = product_name.strip().lower()

    for offset in range(PRODUCTS_START_OFFSET, ROWS_PER_DAY):
        row_num = day_start_row + offset
        cell_product = ws.cell(row=row_num, column=COL_PRODUCT).value

        if cell_product:
            cell_product_normalized = str(cell_product).strip().lower()
            if cell_product_normalized == product_name_normalized:
                return row_num

    return None


def find_empty_row_in_day(ws, day_start_row):
    """
    Najde první prázdný řádek v daném dni pro přidání vlastního produktu.
    Vrací číslo řádku nebo None, pokud není volný řádek.
    Hledá jen nad řádkem "Celkem", aby se nepřepsal součtový vzorec.
    """
    for offset in range(PRODUCTS_START_OFFSET, ROWS_PER_DAY - 1):  # -1 aby zůstal poslední řádek prázdný
        row_num = day_start_row + offset

        cell_number = ws.cell(row=row_num, column=COL_NUMBER).value
        if isinstance(cell_number, str) and cell_number.strip().lower() == 'celkem':
            break

        cell_product = ws.cell(row=row_num, column=COL_PRODUCT).value
        cell_price = ws.cell(row=row_num, column=COL_PRICE).value

        if not cell_product and not cell_price:
            return row_num

    return None


def add_order_to_excel(order_data):
    """
    Přidá objednávku do Excelu pro aktuální den.

    Args:
        order_data (dict): {
            "timestamp": "2024-01-15 14:30:00",
            "items": [
                {"name": "Produkt 1", "price": 100, "quantity": 2},
                {"name": "Produkt 2", "price": 50, "quantity": 1}
            ],
            "total": 250
        }

    Returns:
        dict: {"success": bool, "message": str, "error": str (volitelné)}
    """
    try:
        if not os.path.exists(EXCEL_FILE):
            error_msg = f"Excel soubor nenalezen: {EXCEL_FILE}"
            print(f"CHYBA: {error_msg}")
            return {"success": False, "error": error_msg}

        # Automatická záloha (pokud dnes ještě nebyla)
        auto_backup_if_needed()

        # Otevři Excel
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb[WORKSHEET_NAME]

        # Zjisti aktuální den v měsíci
        timestamp = datetime.strptime(order_data['timestamp'], "%Y-%m-%d %H:%M:%S")
        current_day = timestamp.day

        print(f"\n{'='*60}")
        print(f"Ukládám objednávku pro den {current_day}. {timestamp.month}. {timestamp.year}")
        print(f"{'='*60}")

        # Najdi začátek bloku pro tento den
        day_start_row = get_day_start_row(current_day)
        print(f"Blok pro den {current_day} začíná na řádku {day_start_row}")

        # KONTROLA: Ověř, že tento blok existuje v Excel souboru
        max_row = ws.max_row
        expected_last_row = day_start_row + ROWS_PER_DAY - 1

        if day_start_row > max_row:
            error_msg = (f"Excel soubor má jen {max_row} řádků, "
                        f"ale den {current_day} vyžaduje řádek {day_start_row}. "
                        f"Ujistěte se, že používáte Excel soubor s {current_day} bloky pro tento měsíc.")
            print(f"CHYBA: {error_msg}")
            wb.close()
            return {"success": False, "error": error_msg}

        # Projdi všechny položky v objednávce
        for item in order_data['items']:
            product_name = item['name']
            quantity = item['quantity']
            price = item['price']

            print(f"\nZpracovávám: {product_name} (× {quantity})")

            # Najdi řádek s tímto produktem
            product_row = find_product_row_in_day(ws, day_start_row, product_name)

            if product_row:
                # Produkt nalezen - navýš množství
                current_quantity = ws.cell(row=product_row, column=COL_QUANTITY).value or 0
                new_quantity = current_quantity + quantity
                ws.cell(row=product_row, column=COL_QUANTITY, value=new_quantity)
                print(f"  ✓ Nalezen na řádku {product_row}, množství: {current_quantity} → {new_quantity}")

            else:
                # Produkt nenalezen - najdi prázdný řádek a přidej ho
                empty_row = find_empty_row_in_day(ws, day_start_row)

                if empty_row:
                    ws.cell(row=empty_row, column=COL_PRODUCT, value=product_name)
                    ws.cell(row=empty_row, column=COL_PRICE, value=price)
                    ws.cell(row=empty_row, column=COL_QUANTITY, value=quantity)
                    # Přidej vzorec pro celkovou cenu
                    ws.cell(row=empty_row, column=COL_TOTAL, value=f"=D{empty_row}*C{empty_row}")
                    print(f"  ✓ Nový produkt přidán na řádek {empty_row}")
                else:
                    # Neukládej nic, aby se objednávka nezapsala jen napůl
                    wb.close()
                    error_msg = (f"V Excelu pro den {current_day} už není volný řádek "
                                f"pro zakázku '{product_name}'. Objednávka nebyla uložena.")
                    print(f"  ✗ CHYBA: {error_msg}")
                    return {"success": False, "error": error_msg}

        # Ulož Excel
        wb.save(EXCEL_FILE)
        wb.close()

        print(f"\n{'='*60}")
        print(f"Objednávka úspěšně uložena!")
        print(f"{'='*60}\n")
        return {"success": True, "message": "Objednávka byla úspěšně uložena"}

    except Exception as e:
        error_msg = f"Chyba při ukládání do Excelu: {str(e)}"
        print(error_msg)
        import traceback
        traceback.print_exc()
        return {"success": False, "error": error_msg}


def storno_order_in_excel(order_data):
    """
    Stornuje položky v Excelu pro aktuální den (odečte množství).
    Funguje opačně než add_order_to_excel. Storno je atomické - pokud některou
    položku nelze stornovat (není v Excelu nebo by množství kleslo pod nulu),
    neuloží se nic.

    Args:
        order_data (dict): stejný formát jako u add_order_to_excel

    Returns:
        dict: {"success": bool, "message": str, "error": str (volitelné)}
    """
    try:
        if not os.path.exists(EXCEL_FILE):
            error_msg = f"Excel soubor nenalezen: {EXCEL_FILE}"
            print(f"CHYBA: {error_msg}")
            return {"success": False, "error": error_msg}

        # Automatická záloha (pokud dnes ještě nebyla)
        auto_backup_if_needed()

        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb[WORKSHEET_NAME]

        timestamp = datetime.strptime(order_data['timestamp'], "%Y-%m-%d %H:%M:%S")
        current_day = timestamp.day

        print(f"\n{'='*60}")
        print(f"STORNO pro den {current_day}. {timestamp.month}. {timestamp.year}")
        print(f"{'='*60}")

        day_start_row = get_day_start_row(current_day)

        if day_start_row > ws.max_row:
            error_msg = (f"Excel soubor má jen {ws.max_row} řádků, "
                        f"ale den {current_day} vyžaduje řádek {day_start_row}.")
            print(f"CHYBA: {error_msg}")
            wb.close()
            return {"success": False, "error": error_msg}

        # Sečti množství podle názvu (stejný produkt může být v košíku víckrát)
        totals = {}
        for item in order_data['items']:
            key = item['name'].strip().lower()
            if key not in totals:
                totals[key] = {"name": item['name'].strip(), "quantity": 0}
            totals[key]["quantity"] += item['quantity']

        # 1. Ověř všechny položky dříve, než cokoliv změníš
        changes = []
        errors = []
        for entry in totals.values():
            product_name = entry["name"]
            quantity = entry["quantity"]

            product_row = find_product_row_in_day(ws, day_start_row, product_name)
            if not product_row:
                errors.append(f"'{product_name}' dnes v Excelu není")
                continue

            current_quantity = ws.cell(row=product_row, column=COL_QUANTITY).value or 0
            if current_quantity < quantity:
                errors.append(f"'{product_name}': dnes prodáno jen {current_quantity} ks, nelze stornovat {quantity} ks")
                continue

            changes.append((product_row, product_name, current_quantity, current_quantity - quantity))

        if errors:
            wb.close()
            error_msg = "Storno nelze provést: " + "; ".join(errors)
            print(f"CHYBA: {error_msg}")
            return {"success": False, "error": error_msg}

        # 2. Proveď změny
        for product_row, product_name, current_quantity, new_quantity in changes:
            ws.cell(row=product_row, column=COL_QUANTITY, value=new_quantity)
            print(f"  ✓ {product_name} (řádek {product_row}), množství: {current_quantity} → {new_quantity}")

        wb.save(EXCEL_FILE)
        wb.close()

        print(f"\n{'='*60}")
        print(f"Storno úspěšně uloženo!")
        print(f"{'='*60}\n")
        return {"success": True, "message": "Položky byly stornovány"}

    except Exception as e:
        error_msg = f"Chyba při stornování v Excelu: {str(e)}"
        print(error_msg)
        import traceback
        traceback.print_exc()
        return {"success": False, "error": error_msg}


def get_order_summary_for_day(day):
    """
    Vrátí souhrn objednávek pro daný den (kolik kusů jednotlivých produktů bylo prodáno).
    """
    try:
        if not os.path.exists(EXCEL_FILE):
            return {}

        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb[WORKSHEET_NAME]

        day_start_row = get_day_start_row(day)
        summary = {}

        for offset in range(PRODUCTS_START_OFFSET, ROWS_PER_DAY):
            row_num = day_start_row + offset
            product_name = ws.cell(row=row_num, column=COL_PRODUCT).value
            quantity = ws.cell(row=row_num, column=COL_QUANTITY).value

            if product_name and quantity and quantity > 0:
                summary[product_name] = quantity

        wb.close()
        return summary

    except Exception as e:
        print(f"Chyba při čtení souhrnu: {str(e)}")
        return {}


def backup_excel():
    """
    Vytvoří zálohu Excel souboru s časovým razítkem.
    Zálohy se ukládají do složky 'backup' vedle původního souboru.

    Returns:
        dict: {"success": bool, "backup_path": str, "message": str}
    """
    try:
        if not os.path.exists(EXCEL_FILE):
            return {
                "success": False,
                "message": f"Excel soubor nenalezen: {EXCEL_FILE}"
            }

        # Vytvoř složku pro zálohy
        excel_dir = os.path.dirname(EXCEL_FILE)
        backup_dir = os.path.join(excel_dir, 'backup')

        if not os.path.exists(backup_dir):
            os.makedirs(backup_dir)
            print(f"Vytvořena složka pro zálohy: {backup_dir}")

        # Název souboru se zálohy s časovým razítkem
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        excel_basename = os.path.basename(EXCEL_FILE)
        excel_name, excel_ext = os.path.splitext(excel_basename)
        backup_filename = f"{excel_name}_ZALOHA_{timestamp}{excel_ext}"
        backup_path = os.path.join(backup_dir, backup_filename)

        # Zkopíruj soubor
        shutil.copy2(EXCEL_FILE, backup_path)

        # Zkontroluj, že záloha byla vytvořena
        if os.path.exists(backup_path):
            file_size = os.path.getsize(backup_path) / 1024  # KB
            print(f"✓ Záloha vytvořena: {backup_path} ({file_size:.1f} KB)")

            return {
                "success": True,
                "backup_path": backup_path,
                "backup_filename": backup_filename,
                "message": f"Záloha úspěšně vytvořena: {backup_filename}"
            }
        else:
            return {
                "success": False,
                "message": "Záloha nebyla vytvořena (neznámá chyba)"
            }

    except Exception as e:
        print(f"Chyba při vytváření zálohy: {str(e)}")
        import traceback
        traceback.print_exc()
        return {
            "success": False,
            "message": f"Chyba při zálohování: {str(e)}"
        }


def get_backup_list():
    """
    Vrátí seznam existujících záloh.

    Returns:
        list: Seznam slovníků s informacemi o zálohách
    """
    try:
        excel_dir = os.path.dirname(EXCEL_FILE)
        backup_dir = os.path.join(excel_dir, 'backup')

        if not os.path.exists(backup_dir):
            return []

        backups = []
        for filename in os.listdir(backup_dir):
            if filename.endswith('.xlsx') or filename.endswith('.xls'):
                filepath = os.path.join(backup_dir, filename)
                stat = os.stat(filepath)
                backups.append({
                    'filename': filename,
                    'path': filepath,
                    'size': stat.st_size,
                    'created': datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M:%S")
                })

        # Seřaď podle data vytvoření (nejnovější první)
        backups.sort(key=lambda x: x['created'], reverse=True)
        return backups

    except Exception as e:
        print(f"Chyba při načítání seznamu záloh: {str(e)}")
        return []


def auto_backup_if_needed():
    """
    Automatická záloha, pokud dnes ještě nebyla vytvořena.
    Volá se při prvním uložení objednávky v daný den.

    Returns:
        bool: True pokud byla záloha vytvořena nebo již existuje, False při chybě
    """
    try:
        today = datetime.now().strftime("%Y-%m-%d")
        excel_dir = os.path.dirname(EXCEL_FILE)
        backup_dir = os.path.join(excel_dir, 'backup')

        if not os.path.exists(backup_dir):
            # Pokud složka neexistuje, vytvoř zálohu
            result = backup_excel()
            return result['success']

        # Zkontroluj, zda dnes již byla záloha
        for filename in os.listdir(backup_dir):
            if today in filename:
                print(f"Záloha pro dnešní den již existuje: {filename}")
                return True

        # Dnes ještě nebyla záloha, vytvoř ji
        print(f"Vytvářím automatickou zálohu pro {today}...")
        result = backup_excel()
        return result['success']

    except Exception as e:
        print(f"Chyba při automatické záloze: {str(e)}")
        return False
